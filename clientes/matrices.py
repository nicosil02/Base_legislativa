"""Lectura de las matrices puntuales de PLs que Nicolas mantiene a mano en
Excel (Bayer Crop/Syngenta, Incode) - complementan el filtro por categoria
de scraper/categorias.py::CATEGORIA_CLIENTES_PL, que puede no detectar un PL
realmente relevante si su tema literal no calza con la categoria del cliente
(confirmado en vivo 2026-09-13: 4 PLs reales de la matriz Bayer Crop en Peru
cayeron en categorias "Consumo masivo"/"Tributos", no "Agricultura" - son
leyes tributarias que afectan al agro, no leyes agrarias en si).

Se lee el Excel directo en cada carga (decision de Nicolas 2026-09-13: sigue
editando el archivo como siempre, sin pasos extra) - ambos archivos se
versionan en el repo (ver .gitignore) igual que notas.md, para que la app en
produccion tambien los pueda leer.

Nicolas confirmo 2026-09-13:
- Bayer Crop (hojas PERU y ECUADOR): funciona 100%, usar tal cual.
- Incode Peru (hoja "Peru Legislativo"): desactualizada (PLs archivados al
  pasar al Congreso bicameral) - NO USAR.
- Incode Ecuador (hoja "Ecuador Legislativo"): la lista de PLs es real, solo
  el estado que trae la matriz puede estar desactualizado - se usa el
  numero de PL para cruzar, no el estado_matriz como fuente de verdad.
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

CLIENTES_DIR = Path(__file__).resolve().parent

# La columna "Proyecto" de la matriz Peru trae texto libre alrededor del
# codigo real ("N° 00011-2026-2031-CD", "Proposición N° 00080-...") -
# se extrae solo el codigo (mismo formato que proyecto_ley en nuestra DB:
# 5 digitos-4 digitos-4 digitos-2/3 letras) para poder cruzar por igualdad
# exacta contra pages/1_Peru.py::df["PL"].
_RE_PL_PE = re.compile(r"\d+-\d{4}-\d{4}-[A-Za-z]{1,3}")


def _normalizar_pl_pe(texto: str) -> str | None:
    m = _RE_PL_PE.search(str(texto))
    return m.group(0).upper() if m else None

BAYER_CROP_XLSX = CLIENTES_DIR / "MatrizPLAndinos_Bayer CROP.xlsx"
INCODE_XLSX = CLIENTES_DIR / "140125 - Matriz - Matriz Legislativa y Regulatoria INCODE.xlsx"


_RE_DIGITS = re.compile(r"\d{5,}")


def _split_numeros(valor, extraer_digitos: bool = False) -> list[str]:
    """Una celda puede traer un solo numero o varios separados por salto de
    linea/coma/barra (PLs unificados en Ecuador, o formato mixto tipo
    "Cod. AN-2024-2928 / 450889") - normaliza a lista de strings.

    `extraer_digitos`: para Ecuador, si un token no es puramente numerico
    (trae el "codigo de tramite" AN-XXX en vez del n_tramite numerico que
    usamos para cruzar), se intenta extraer la corrida de 5+ digitos que
    hay adentro. NO usar para Peru - ahi el codigo real ES alfanumerico
    ("00011-2026-2031-CD") y este atajo lo rompería."""
    if valor is None:
        return []
    if isinstance(valor, (int, float)):
        return [str(int(valor))]
    texto = str(valor).replace(",", "\n").replace("/", "\n")
    tokens = [p.strip() for p in texto.split("\n") if p.strip()]
    if not extraer_digitos:
        return tokens
    out = []
    for t in tokens:
        m = _RE_DIGITS.search(t)
        out.append(m.group(0) if m else t)
    return out


def _leer_hoja(path: Path, hoja: str, header_row: int,
                cols: dict[str, str | list[str]]) -> list[dict]:
    """Lee una hoja mapeando nombres de columna (`cols`: alias -> header real,
    o lista de variantes aceptadas - ej. la hoja ECUADOR tiene el typo
    "Tit�lo" en vez de "T�tulo") a partir de la fila de headers. Filas
    totalmente vacias se saltean. Si el archivo/hoja no existe, devuelve
    lista vacia (no rompe la app si Nicolas todavia no subio o renombro algo)."""
    import openpyxl

    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    if hoja not in wb.sheetnames:
        return []
    ws = wb[hoja]
    # Los headers reales tienen inconsistencias (espacios de mas, typos como
    # "Tit�lo" en vez de "T�tulo") - matchear normalizado (strip + lower) en
    # vez de comparacion exacta.
    headers_norm = [
        (str(c.value).strip().lower() if c.value is not None else None)
        for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]
    idx = {}
    for alias, candidatos in cols.items():
        if isinstance(candidatos, str):
            candidatos = [candidatos]
        for real in candidatos:
            real_norm = real.strip().lower()
            if real_norm in headers_norm:
                idx[alias] = headers_norm.index(real_norm)
                break
    out = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c is not None for c in row):
            continue
        out.append({alias: row[i] for alias, i in idx.items()})
    return out


def matriz_bayer_crop(pais: str) -> list[dict]:
    """PLs puntuales de la matriz Bayer Crop/Syngenta para `pais` ("PE"/"EC").
    Cada dict: pl_numero, titulo_matriz, tema_matriz, estado_matriz."""
    hoja = "PERÚ" if pais == "PE" else "ECUADOR"
    filas = _leer_hoja(
        BAYER_CROP_XLSX, hoja, header_row=(3 if pais == "PE" else 2),
        cols={"titulo": ["Título", "Titúlo"], "proyecto": "Proyecto", "tema": "Tema", "estado": "Estado"},
    )
    out = []
    for f in filas:
        for num in _split_numeros(f.get("proyecto"), extraer_digitos=(pais == "EC")):
            if pais == "PE":
                num = _normalizar_pl_pe(num)
                if not num:
                    continue
            out.append({
                "pl_numero": num, "titulo_matriz": f.get("titulo"),
                "tema_matriz": f.get("tema"), "estado_matriz": f.get("estado"),
            })
    return out


def matriz_incode_ec() -> list[dict]:
    """PLs puntuales de la matriz Incode para Ecuador (la unica hoja que
    Nicolas confirmo vigente - la de Peru esta desactualizada, no se lee)."""
    filas = _leer_hoja(
        INCODE_XLSX, "Ecuador Legislativo", header_row=2,
        cols={"titulo": ["Título", "Titúlo"], "proyecto": "Proyecto", "tema": "Tema", "estado": "Estado"},
    )
    out = []
    for f in filas:
        for num in _split_numeros(f.get("proyecto"), extraer_digitos=True):
            out.append({
                "pl_numero": num, "titulo_matriz": f.get("titulo"),
                "tema_matriz": f.get("tema"), "estado_matriz": f.get("estado"),
            })
    return out


# ============================================================
# Matching noticia <-> PL trackeado por PARTE DEL NOMBRE, no por numero
# ============================================================
# Nicolas confirmo 2026-09-13: el equipo identifica un PL por parte de su
# nombre/titulo (ej. "Ley de Paramos", "Ley del INIAP"), nunca por el numero
# de tramite - confirma el hallazgo empirico (0 de 2145 noticias EC citan un
# numero de tramite exacto). Esto es la pieza real de la Fase 3: detectar
# cuando una noticia (que se actualiza rapido) menciona el nombre de un PL
# que ya trackeamos (matriz), como señal temprana antes que el portal.

# Boilerplate legislativo/generico que NO sirve para identificar un PL
# puntual (aparece en casi todos los titulos por igual) - se saca antes de
# buscar coincidencias, para no matchear por "proyecto de ley organica".
_STOPWORDS_PL = {
    "proyecto", "proyectos", "ley", "leyes", "organica", "organico",
    "reformatoria", "reformatorio", "reforma", "codigo", "para", "que",
    "sobre", "con", "del", "las", "los", "una", "uno", "por", "sus", "en",
    "de", "la", "el", "y", "a", "al", "se", "su", "e", "u", "o", "no", "es",
    "un", "lo", "como", "mas", "integral", "nacional", "ecuador", "asamblea",
    "diversos", "varios", "articulos", "articulo", "cuerpos", "legales",
    "vigente", "vigentes", "cumplimiento", "fin", "efectos", "materia",
}


# Palabras genericas que aparecen en MUCHOS titulos de PL sin identificar
# nada puntual (confirmado en vivo 2026-09-13: "desarrollo" solo hizo
# matchear un PL de desarrollo agropecuario contra 1444 noticias sin
# relacion real - cripto, deportes, narcotrafico). Se suman al stopword
# list angosto de arriba para las palabras "normales" (no siglas).
_PALABRAS_GENERICAS_PL = {
    "desarrollo", "unificado", "productivo", "productiva", "general",
    "publico", "publica", "sistema", "gestion", "fortalecimiento",
    "proteccion", "regimen", "responsabilidad", "seguridad", "servicio",
    "servicios", "derecho", "derechos", "modelo", "plan", "politica",
    "sancionar", "sanciones", "control", "registro", "personas",
}

# Siglas/instituciones en MAYUSCULAS dentro del titulo ORIGINAL (antes de
# normalizar) - "INIAP", "COIP", "OVM" identifican un PL puntual por si
# solas, a diferencia de una palabra comun del idioma. Denylist corta para
# los falsos positivos reales de este corpus (titulares con el pais en
# mayusculas como dateline, ej. "ECUADOR: Comision aprueba...").
_RE_ACRONIMO = re.compile(r"\b[A-Z]{3,}\b")
_ACRONIMOS_IGNORAR = {"ECUADOR", "PERU", "LEY", "ASAMBLEA", "COMISION", "PROYECTO"}


def _normalizar_texto(s: str | None) -> str:
    s = (s or "").lower()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _acronimos(texto: str | None) -> set[str]:
    return {a for a in _RE_ACRONIMO.findall(texto or "") if a not in _ACRONIMOS_IGNORAR}


def palabras_clave_titulo(titulo: str | None, min_len: int = 5) -> set[str]:
    """Palabras distintivas de un titulo (sin el boilerplate legislativo
    generico ni las palabras demasiado comunes) - lo que realmente
    identifica a ESE PL puntual, sin contar las siglas (ver _acronimos)."""
    texto = _normalizar_texto(titulo)
    palabras = re.findall(r"[a-z]+", texto)
    return {
        p for p in palabras
        if len(p) >= min_len and p not in _STOPWORDS_PL and p not in _PALABRAS_GENERICAS_PL
    }


# Los PLs trackeados son todos de Ecuador (ver pls_trackeados_ec) - si una
# noticia nombra OTRO pais y nunca menciona Ecuador, el match de la via 2
# (palabras compartidas) es casi siempre casualidad de vocabulario, no la
# misma noticia. Caso real confirmado en vivo 2026-09-14: "Congreso de Perú
# y Meta evalúan cooperación en inteligencia artificial y seguridad digital"
# matcheaba un PL ecuatoriano de IA solo por compartir "inteligencia",
# "artificial", "digital" - ninguna de esas palabras identifica al PL, son
# genericas del tema. Mismo problema con notas de México, España y T-MEC/
# Norteamerica encontradas en la misma auditoria. Subir minimo_palabras no
# sirve (ya probado: mata matches reales como "... obliga a Ecuador a
# decidir ahora" antes de sacar los falsos positivos). Solo se aplica a la
# via 2 (palabras) - una sigla compartida (INIAP, COIP...) es una senal
# fuerte por si sola, muy improbable que otro pais tenga la misma sigla.
_OTRO_PAIS_MARCADORES = {
    "peru", "mexico", "colombia", "chile", "argentina", "espana",
    "bolivia", "venezuela", "brasil", "uruguay", "paraguay", "norteamerica",
    "euros",  # Ecuador usa USD - "euros" es senal fuerte de nota europea
}
_ECUADOR_MARCADORES = {
    "ecuador", "ecuatoriano", "ecuatoriana", "ecuatorianos", "ecuatorianas",
    "quito", "guayaquil",
}


def _parece_de_otro_pais(texto_normalizado: str) -> bool:
    palabras = set(re.findall(r"[a-z]+", texto_normalizado))
    if palabras & _ECUADOR_MARCADORES:
        return False
    return bool(palabras & _OTRO_PAIS_MARCADORES)


def coincide_con_noticia(pl_titulo: str | None, noticia_texto: str | None, minimo_palabras: int = 3,
                          noticia_titulo: str | None = None) -> bool:
    """True si `noticia_texto` (titulo+resumen de una noticia) parece hablar
    del mismo PL. Dos vias, ambas deterministas/auditables (no es matching
    difuso tipo TF-IDF, ya descartado en esta sesion por impreciso):
    1. Comparten una SIGLA real (INIAP, COIP, OVM...) - identificador fuerte
       por si solo, confirmado con el caso real del PL del INIAP.
    2. Comparten `minimo_palabras` (3 por defecto) palabras distintivas del
       titulo - probado en vivo contra las 2145 noticias EC reales: con 2
       palabras salian 248 matches (ej. "acceso"+"recursos" emparejando una
       noticia de fintech mexicano sin relacion real); con 3 bajo a 18,
       manteniendo los aciertos reales (INIAP, IA). Ademas exige que la
       noticia no parezca ser de otro pais (ver _parece_de_otro_pais). Sigue
       siendo una señal para que Nicolas revise, no una alerta automatica -
       puede quedar algun falso positivo, pero ya en un volumen chico y
       revisable."""
    acr_pl = _acronimos(pl_titulo)
    if acr_pl and (acr_pl & _acronimos(noticia_texto)):
        return True
    kw_pl = palabras_clave_titulo(pl_titulo)
    if len(kw_pl) < minimo_palabras:
        return False
    kw_noticia = palabras_clave_titulo(noticia_texto, min_len=4)
    if len(kw_pl & kw_noticia) < minimo_palabras:
        return False
    # El chequeo de "otro pais" solo mira el TITULO (no el resumen completo)
    # - un pais nombrado de paso en el cuerpo (ej. "fondos desde Venezuela"
    # en una nota ecuatoriana de lavado de activos) no debe descartar un
    # match real; que el pais aparezca en el titular es señal mucho mas
    # fuerte de que la noticia ES sobre ese pais. Si no se pasa el titulo
    # por separado, cae al texto completo (comportamiento previo).
    texto_pais = noticia_titulo if noticia_titulo is not None else noticia_texto
    return not _parece_de_otro_pais(_normalizar_texto(texto_pais))


def pls_trackeados_ec() -> list[dict]:
    """Todos los PLs EC trackeados (Bayer Crop/Syngenta + Incode), con
    `clientes` agregado - para chequear noticias contra el set completo de
    una sola vez en vez de repetir la lectura del Excel por cliente."""
    out = []
    for f in matriz_bayer_crop("EC"):
        out.append({**f, "clientes": ["bayer", "syngenta"]})
    for f in matriz_incode_ec():
        out.append({**f, "clientes": ["incode"]})
    return out


# ============================================================
# self-check (Ponytail: 1 chequeo ejecutable de la logica no trivial)
# ============================================================

def _demo():
    pe = matriz_bayer_crop("PE")
    ec = matriz_bayer_crop("EC")
    inc = matriz_incode_ec()
    assert isinstance(pe, list) and isinstance(ec, list) and isinstance(inc, list)
    if BAYER_CROP_XLSX.exists():
        assert len(pe) > 0, "esperaba PLs reales en la matriz Bayer Crop PE"
        assert len(ec) > 0, "esperaba PLs reales en la matriz Bayer Crop EC"
        assert all(f["pl_numero"] and f["titulo_matriz"] for f in pe), (
            "algun PL de PE quedo sin titulo - revisar header 'Titulo/Titulo'")
        assert all(f["pl_numero"] and f["titulo_matriz"] for f in ec), (
            "algun PL de EC quedo sin titulo - la hoja ECUADOR tiene el typo 'Titulo'")
        print(f"OK matriz_bayer_crop('PE'): {len(pe)} PLs")
        print(f"OK matriz_bayer_crop('EC'): {len(ec)} PLs")
    else:
        print("SKIP: no se encontro el Excel de Bayer Crop en este entorno")
    if INCODE_XLSX.exists():
        assert len(inc) > 0, "esperaba PLs reales en la matriz Incode Ecuador"
        print(f"OK matriz_incode_ec(): {len(inc)} PLs")
    else:
        print("SKIP: no se encontro el Excel de Incode en este entorno")

    # Matching por nombre - caso real verificado en vivo 2026-09-13: una
    # noticia real sobre el INIAP matcheo el titulo de ese PL puntual.
    pl_iniap = "Proyecto de ley reformatoria a la Ley Constitutiva del INIAP"
    noticia_real = "ECUADOR: Comisión aprueba informe para reformar la Ley del INIAP"
    noticia_no_relacionada = "Presidenta Fujimori entrega ayuda humanitaria a población de Purús"
    assert coincide_con_noticia(pl_iniap, noticia_real), (
        "deberia matchear - caso real verificado (INIAP)")
    assert not coincide_con_noticia(pl_iniap, noticia_no_relacionada), (
        "no deberia matchear - sin relacion real"
    )
    print("OK coincide_con_noticia: matchea el caso real (INIAP), no matchea uno sin relacion")

    trackeados = pls_trackeados_ec()
    if BAYER_CROP_XLSX.exists() or INCODE_XLSX.exists():
        assert len(trackeados) == len(ec) + len(inc)
        assert all("clientes" in f for f in trackeados)
        print(f"OK pls_trackeados_ec(): {len(trackeados)} PLs con cliente asignado")


if __name__ == "__main__":
    _demo()
